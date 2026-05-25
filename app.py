import streamlit as st
import pandas as pd
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm
import io
import os
import PyPDF2
import google.generativeai as genai
import json

# --- PAGE CONFIGURATION ---
st.set_page_config(page_title="Weekly Report Automator", page_icon="📄", layout="wide")

st.markdown("""
    <style>
        [data-testid="stAppViewContainer"] { background-color: #f8fafc; }
        h1, h2, h3, h4 { font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; font-weight: 800; }
        [data-testid="stFileUploadDropzone"] { border: 2px dashed #3b82f6; border-radius: 12px; padding: 20px; background-color: rgba(59, 130, 246, 0.05); }
        [data-testid="baseButton-primary"] { background-color: #3b82f6; color: white; border-radius: 8px; font-weight: bold; border: none; transition: all 0.3s ease; }
        [data-testid="baseButton-primary"]:hover { background-color: #2563eb; box-shadow: 0 4px 10px rgba(59, 130, 246, 0.3); }
    </style>
""", unsafe_allow_html=True)

st.title("📄 Weekly Report Automator (AI-Powered)")
st.divider()

# --- INITIALIZE SESSION STATE FOR AI TEXT ---
if "ai_benefits" not in st.session_state: st.session_state.ai_benefits = ""
if "ai_objectives" not in st.session_state: st.session_state.ai_objectives = ""
if "ai_outcomes" not in st.session_state: st.session_state.ai_outcomes = ""
if "ai_outline" not in st.session_state:
    # Default empty Pandas structure
    st.session_state.ai_outline = pd.DataFrame([{"Topic": "", "Sub-topics": "", "Activities": ""}])

# --- AI EXTRACTION SECTION ---
st.subheader("🤖 Step 1: AI Module Extraction")
api_key = st.text_input("Enter your Google Gemini API Key", type="password")
uploaded_module = st.file_uploader("Upload Instructor's Module (PDF only for now)", type=["pdf"])

if st.button("✨ Extract Course Details with AI", type="primary"):
    if not api_key:
        st.error("⚠️ Please enter your Gemini API Key.")
    elif not uploaded_module:
        st.error("⚠️ Please upload a PDF module.")
    else:
        with st.spinner("AI is reading the module and extracting data..."):
            try:
                pdf_reader = PyPDF2.PdfReader(uploaded_module)
                raw_text = ""
                for page in pdf_reader.pages:
                    raw_text += page.extract_text() + "\n"

                genai.configure(api_key=api_key)
                model = genai.GenerativeModel('gemini-2.5-flash')

                prompt = f"""
                You are an expert educational curriculum extractor. Read the following course module text.
                Extract the information and return it EXACTLY as a raw JSON dictionary with NO markdown formatting.

                The JSON must have these exact keys:
                "benefits": "A short summary of the module benefits."
                "learning_objectives": "A short paragraph of the learning objectives."
                "learning_outcomes": "A short paragraph of the learning outcomes."
                "outline": [ {{"Topic": "Main Topic 1", "Sub-topics": "Subtopic details", "Activities": "Suggested class activity"}}, ... ]

                Course Text:
                {raw_text[:15000]}
                """

                response = model.generate_content(prompt)
                cleaned_response = response.text.replace("```json", "").replace("```", "").strip()
                extracted_data = json.loads(cleaned_response)

                st.session_state.ai_benefits = extracted_data.get("benefits", "")
                st.session_state.ai_objectives = extracted_data.get("learning_objectives", "")
                st.session_state.ai_outcomes = extracted_data.get("learning_outcomes", "")

                # Convert AI outline list into a Pandas DataFrame!
                outline_data = extracted_data.get("outline", [{"Topic": "", "Sub-topics": "", "Activities": ""}])
                st.session_state.ai_outline = pd.DataFrame(outline_data)

                st.success("✅ AI Extraction Complete! Scroll down to review and edit.")
            except Exception as e:
                st.error(f"AI Extraction Failed: {e}")

st.divider()

# --- THE INPUT FORM ---
with st.form("report_form"):
    st.subheader("Step 2: General Information")
    col1, col2, col3 = st.columns(3)
    with col1:
        week_number = st.text_input("Week Number")
        reporting_period = st.text_input("Reporting Period")
        second_admin = st.text_input("Second Admin Name (Optional)")
    with col2:
        submission_date = st.text_input("Submission Date")
        trainer_name = st.text_input("Trainer Name")
    with col3:
        module_number = st.text_input("Module Number")
        module_title = st.text_input("Module Title")

    st.subheader("Step 3: AI-Assisted Learning Outcomes")
    benefits = st.text_area("Benefits of Module", value=st.session_state.ai_benefits, height=100)

    # --- PANDAS UI FOR OUTLINE TABLE ---
    st.markdown("##### Detailed Outline (Interactive Table)")
    st.info("Edit the AI's outline directly in this table before generating the document.")
    outline_df = st.data_editor(st.session_state.ai_outline, num_rows="dynamic", use_container_width=True)

    learning_objectives = st.text_area("Learning Objectives", value=st.session_state.ai_objectives, height=100)
    learning_outcomes = st.text_area("Learning Outcomes", value=st.session_state.ai_outcomes, height=100)

    knowledge_gained = st.text_area("Knowledge Gained")
    competencies_achieved = st.text_area("Competencies Achieved")
    demonstrated_skills = st.text_area("Demonstrated Skills")

    st.subheader("Step 4: Metrics & Assessment")
    col_a, col_b, col_c = st.columns(3)
    with col_a: attendance = st.text_input("Attendance / Punctuality")
    with col_b: engagement_level = st.text_input("Engagement Level")
    with col_c: participant_feedback = st.text_input("Participant Feedback")

    col_q1, col_q2 = st.columns(2)
    with col_q1:
        content_quality = st.text_input("Content Quality")
        materials_resources = st.text_input("Materials & Resources")
        logistical_arrangements = st.text_input("Logistical Arrangements")
    with col_q2:
        trainer_effectiveness = st.text_input("Trainer Effectiveness")
        technical_support = st.text_input("Technical Support")

    observations = st.text_area("Observations", height=100)

    st.subheader("Step 5: Action Items")
    default_actions = pd.DataFrame([
        {"Category": "Immediate Actions", "Action": "", "Responsibility": "", "Deadline": "", "Status": ""},
        {"Category": "Recommendations", "Action": "", "Responsibility": "", "Deadline": "", "Status": ""},
        {"Category": "Follow-Up", "Action": "", "Responsibility": "", "Deadline": "", "Status": ""},
        {"Category": "Resource Improvement", "Action": "", "Responsibility": "", "Deadline": "", "Status": ""},
        {"Category": "Content Modification", "Action": "", "Responsibility": "", "Deadline": "", "Status": ""},
    ])
    action_df = st.data_editor(default_actions, num_rows="dynamic", use_container_width=True)

    st.subheader("Step 6: Photo Gallery")
    uploaded_images = st.file_uploader("Upload exactly 5 Class Photos", type=["jpg", "jpeg", "png"],
                                       accept_multiple_files=True)

    col_p1, col_p2 = st.columns(2)
    with col_p1:
        pic_1_desc = st.text_input("Photo 1 Description")
        pic_2_desc = st.text_input("Photo 2 Description")
        pic_3_desc = st.text_input("Photo 3 Description")
    with col_p2:
        pic_4_desc = st.text_input("Photo 4 Description")
        pic_5_desc = st.text_input("Photo 5 Description")

    st.markdown("<br>", unsafe_allow_html=True)
    submit_button = st.form_submit_button("🚀 Generate Final Word Document", type="primary", use_container_width=True)


def get_safe_action(df, category_name):
    row = df[df["Category"] == category_name]
    if not row.empty: return str(row.iloc[0]["Action"])
    return ""


def get_safe_meta(df):
    if not df.empty: return str(df.iloc[0]["Responsibility"]), str(df.iloc[0]["Deadline"]), str(df.iloc[0]["Status"])
    return "", "", ""


if submit_button:
    if len(uploaded_images) != 5:
        st.error(f"⚠️ Please upload exactly 5 images.")
    elif not os.path.exists("Master_Template.docx"):
        st.error("⚠️ 'Master_Template.docx' not found!")
    else:
        with st.spinner("Building report..."):
            try:
                doc = DocxTemplate("Master_Template.docx")

                image_context = {}
                for i, img in enumerate(uploaded_images):
                    img_bytes = io.BytesIO(img.getvalue())
                    image_context[f"PICTURE_{i + 1}"] = InlineImage(doc, img_bytes, width=Mm(150))

                # --- DYNAMIC PYTHON TABLE DRAWING ---
                outline_subdoc = doc.new_subdoc()

                # Convert the Pandas DataFrame back into a list of rows
                table_records = outline_df.to_dict('records')

                # Draw table based on how many rows Pandas has
                table = outline_subdoc.add_table(rows=1, cols=3)

                # MAGIC FIX: Steal the border style from the first table in your document!
                if len(doc.tables) > 0:
                    table.style = doc.tables[0].style

                    # Headers
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = 'Topic'
                hdr_cells[1].text = 'Sub-topics'
                hdr_cells[2].text = 'Activities'

                # Populate from the Edited Pandas UI
                for item in table_records:
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(item.get("Topic", ""))
                    row_cells[1].text = str(item.get("Sub-topics", ""))
                    row_cells[2].text = str(item.get("Activities", ""))

                # --- EXTRACT ACTIONS ---
                immediate = get_safe_action(action_df, "Immediate Actions")
                recommend = get_safe_action(action_df, "Recommendations")
                follow_up = get_safe_action(action_df, "Follow-Up")
                improvement = get_safe_action(action_df, "Resource Improvement")
                content_mod = get_safe_action(action_df, "Content Modification")
                resp, dead, stat = get_safe_meta(action_df)

                admin_text = f" & {second_admin}" if second_admin.strip() else ""

                context = {
                    "WEEK_NUMBER": week_number, "REPORTING_PERIOD": reporting_period,
                    "SUBMISSION_DATE": submission_date,
                    "TRAINER_NAME": trainer_name, "MODULE_NUMBER": module_number, "MODULE_TITLE": module_title,
                    "SECOND_ADMIN": admin_text,
                    "benefits": benefits, "learning_objectives": learning_objectives,
                    "learning_outcomes": learning_outcomes,

                    "OUTLINE_TABLE": outline_subdoc,

                    "knowledge_gained": knowledge_gained, "compentencies_achieved": competencies_achieved,
                    "demonstrated_skills": demonstrated_skills,
                    "attendance": attendance, "engagement_level": engagement_level,
                    "participant_feedback": participant_feedback,
                    "content_quality": content_quality, "trainer_effectiveness": trainer_effectiveness,
                    "materials_resources": materials_resources,
                    "technical_support": technical_support, "logistical_arrangements": logistical_arrangements,
                    "Observations": observations,
                    "IMMEDIATE": immediate, "RECOMMEND": recommend, "FOLLOW_UP": follow_up, "IMPROVEMENT": improvement,
                    "CONTENT": content_mod, "RESPONSIBILITY": resp, "DEADLINE": dead, "STATUS": stat,
                    "Picture_1_Description": pic_1_desc, "Picture_2_Description": pic_2_desc,
                    "Picture_3_Description": pic_3_desc,
                    "Picture_4_Description": pic_4_desc, "Picture_5_Description": pic_5_desc,
                }

                context.update(image_context)
                doc.render(context)

                output = io.BytesIO()
                doc.save(output)
                output.seek(0)

                st.success("✨ Success!")
                st.download_button("📥 Download Final Report (.docx)", data=output, file_name=f"Weekly_Report.docx",
                                   mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                   type="primary", use_container_width=True)

            except Exception as e:
                st.error(f"Error: {e}")